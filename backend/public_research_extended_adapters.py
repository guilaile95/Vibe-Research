"""Bounded native market/company adapters; source bytes remain immutable evidence.

Source contracts adapted from pinned Vibe AStock d5bf80d, Apache-2.0; see
canonical docs/third-party notice. No upstream Python/JavaScript is executed.
"""
from datetime import date, datetime, timedelta, timezone
from html.parser import HTMLParser
from io import BytesIO
import math
import re
import zipfile
from zoneinfo import ZoneInfo

import public_research_contract as C
import public_research_extended_contract as E
import trade_calendar
from reporting_currency import reporting_currency


def num(value, *, lo=-1e18, hi=1e18):
    if value in (None, '', '--', '-'):
        return None
    if type(value) is str and re.fullmatch(r'-?\d+(?:\.\d+)?', value):
        value = float(value)
    if type(value) not in (float, int) or not math.isfinite(value) or not lo <= value <= hi:
        raise ValueError('invalid source-specific number')
    return float(value)


def obj(value):
    if type(value) is not dict:
        raise ValueError('provider object required')
    return value


def source_day(value):
    # Explicit datacenter midnight timestamp or calendar date; arbitrary suffixes rejected.
    if type(value) is str and re.fullmatch(r'\d{4}-\d{2}-\d{2} 00:00:00', value):
        value = value[:10]
    return C.day(value)


def calendar(cutoff):
    sessions = trade_calendar._load_calendar()
    if not sessions or cutoff not in sessions:
        raise ValueError('cutoff not confirmed by offline CN exchange calendar')
    return [d for d in sessions if d <= cutoff]


def _window(points, expected, field):
    gaps = [d for d in expected if d not in points or points[d].get(field) is None]
    return {'expected_dates': expected, 'observed_dates': [d for d in expected if d not in gaps], 'missing_dates': gaps}


def index_trend(payload, key, cutoff, captured):
    data = obj(obj(payload).get('data'))
    if data.get('code') != key[2:] or type(data.get('market')) is not int or data['market'] != (1 if key.startswith('sh') else 0):
        raise ValueError('index identity mismatch')
    points = {}
    for line in C.rows(data.get('klines')):
        if type(line) is not str or len(line) > 256:
            raise ValueError('invalid index bar')
        parts = line.split(',')
        if len(parts) != 3:
            raise ValueError('unexpected index fields')
        day = C.day(parts[0])
        if day > cutoff:
            continue  # Request cutoff is explicit, retained count below.
        close, amount = num(parts[1], lo=0), num(parts[2], lo=0)
        if close is None or close <= 0:
            raise ValueError('nonpositive close')
        row = {'date': day, 'close': close, 'amount': amount}
        if day in points and points[day] != row:
            raise ValueError('conflicting same-day bar')
        points[day] = row
    sessions = calendar(cutoff)
    if any(d not in sessions for d in points):
        raise ValueError('non-session source date')
    close = points.get(cutoff, {}).get('close')
    metrics, windows = {}, {}
    for n in (5, 20, 60, 250):
        expected = sessions[-n-1:]
        coverage = _window(points, expected, 'close'); windows[str(n)] = coverage
        metrics[f'return_{n}d_pct'] = (close / points[expected[0]]['close'] - 1) * 100 if len(expected) == n+1 and not coverage['missing_dates'] and close else None
        if n in (20, 60, 250):
            dates = sessions[-n:]; cover = _window(points, dates, 'close')
            metrics[f'ma{n}'] = sum(points[d]['close'] for d in dates) / n if len(dates) == n and not cover['missing_dates'] else None
    prior = [d for d in sessions if d[:4] < cutoff[:4]]
    ytd_dates = [d for d in sessions if d >= prior[-1]] if prior else []
    ytd = _window(points, ytd_dates, 'close')
    metrics['ytd_pct'] = (close / points[prior[-1]]['close'] - 1) * 100 if prior and not ytd['missing_dates'] and close else None
    metrics['ytd_basis_date'] = prior[-1] if prior else None
    win = sessions[-250:]
    complete250 = len(win) == 250 and not _window(points, win, 'close')['missing_dates']
    metrics['high250'] = max(points[d]['close'] for d in win) if complete250 else None
    metrics['low250'] = min(points[d]['close'] for d in win) if complete250 else None
    metrics['off_high250_pct'] = (close / metrics['high250'] - 1) * 100 if complete250 and close else None
    local_capture = C.stamp(captured).astimezone(ZoneInfo('Asia/Shanghai'))
    provisional = cutoff == local_capture.date().isoformat() and (local_capture.hour, local_capture.minute) < (15, 0)
    return [{'id': key, 'index_code': key[2:], 'exchange': 'SSE' if key.startswith('sh') else 'BSE' if key.startswith('bj') else 'SZSE',
             'name': E.INDICES[key], 'date': cutoff, 'close': close, 'currency': None, 'unit': 'index points',
             'amount_currency': 'CNY', 'amount_unit': 'CNY', 'adjustment': 'UNADJUSTED',
             'source_vintage': 'SINGLE_RAW_RECEIPT', 'provisional': provisional, 'metrics': metrics,
             'windows': windows, 'ytd_window': ytd, 'amount_window': _window(points, sessions[-60:], 'amount'),
             'amount_rows': [points[d] for d in sessions[-60:] if d in points],
             'missing': ['PROVISIONAL_CURRENT_SESSION'] if provisional else []}], []


class Tables(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True); self.tables=[]; self.table=None; self.row=None; self.cell=None; self.title=''; self.in_title=False; self.nodes=0; self.heading=''; self.in_heading=False; self.contexts=[]; self.table_context=''
    def handle_starttag(self, tag, attrs):
        self.nodes += 1
        if self.nodes > 30000: raise ValueError('HTML node budget')
        if tag == 'title': self.in_title=True
        if tag in ('h1','h2','h3','h4'): self.heading=''; self.in_heading=True
        if tag == 'caption': self.in_heading=True
        if tag == 'table': self.table=[]; self.table_context=self.heading
        if tag == 'tr' and self.table is not None: self.row=[]
        if tag in ('td', 'th') and self.row is not None: self.cell=''
    def handle_data(self, data):
        if self.in_heading:
            self.heading += data
            if len(self.heading)>1024: raise ValueError('HTML heading budget')
        if self.cell is not None:
            self.cell += data
            if len(self.cell)>4096: raise ValueError('HTML cell budget')
        if self.in_title:
            self.title += data
            if len(self.title)>1024: raise ValueError('HTML title budget')
    def handle_endtag(self, tag):
        if tag == 'title': self.in_title=False
        if tag in ('h1','h2','h3','h4','caption'):
            self.in_heading=False
            if tag=='caption': self.table_context=self.heading
        if tag in ('td','th') and self.cell is not None:
            self.row.append(self.cell.strip()); self.cell=None
        if tag == 'tr' and self.row is not None:
            self.table.append(self.row); self.row=None
        if tag == 'table' and self.table is not None:
            self.tables.append(self.table); self.contexts.append(self.table_context); self.table=None


def treasury(raw, cutoff):
    parser=Tables(); parser.feed(raw.decode('utf-8'))
    header=['曲线名称','日期','3月','6月','1年','3年','5年','7年','10年','30年']
    tables=[t for t in parser.tables if t and t[0]==header]
    if len(tables)!=1: raise ValueError('treasury header drift')
    output=[]
    for row in C.rows(tables[0][1:]):
        if len(row)!=10 or row[0]!='中债国债收益率曲线' or C.day(row[1])!=cutoff: raise ValueError('treasury identity/date mismatch')
        for tenor,value in zip(('3m','6m','1y','3y','5y','7y','10y','30y'),row[2:]):
            output.append({'id': cutoff+':'+tenor,'date':cutoff,'curve':'treasury','tenor':tenor,'yield_pct':num(value,lo=-100,hi=100),
                           'currency':'CNY','unit':'percent yield','price':None,'holding_return':None})
    if len({r['id'] for r in output})!=len(output): raise ValueError('duplicate treasury curve')
    return output, ['BOND_CALENDAR_COVERAGE_NOT_PROVEN']


def excel_rows(raw):
    import xlrd
    from openpyxl.utils.exceptions import InvalidFileException
    from xml.etree.ElementTree import ParseError
    import struct
    try:
        return _excel_rows(raw)
    except (xlrd.biffh.XLRDError, zipfile.BadZipFile, InvalidFileException, ParseError, struct.error, EOFError, IndexError) as exc:
        raise ValueError('malformed source spreadsheet') from exc


def _excel_rows(raw):
    if raw[:2]==b'PK':
        with zipfile.ZipFile(BytesIO(raw)) as archive:
            infos=archive.infolist()
            if len(infos)>64 or sum(i.file_size for i in infos)>8*1024*1024 or any(i.file_size>4*1024*1024 for i in infos): raise ValueError('XLSX expansion budget')
        import openpyxl
        book=openpyxl.load_workbook(BytesIO(raw),read_only=True,data_only=True,keep_links=False)
        try:
            if len(book.worksheets)!=1: raise ValueError('unexpected worksheets')
            sheet=book.worksheets[0]
            if sheet.max_row is None or sheet.max_column is None or sheet.max_row>1000 or sheet.max_column>32: raise ValueError('spreadsheet cell budget')
            return list(sheet.values)
        finally: book.close()
    import xlrd
    book=xlrd.open_workbook(file_contents=raw,on_demand=True)
    try:
        if book.nsheets!=1: raise ValueError('unexpected worksheets')
        # Preflight BIFF records before xlrd allocates sparse worksheet cells.
        import struct
        position=book._sh_abs_posn[0]; records=0
        while position+4<=len(book.mem):
            code,length=struct.unpack_from('<HH',book.mem,position); position+=4; records+=1
            if records>40000 or position+length>len(book.mem): raise ValueError('BIFF record budget')
            record=book.mem[position:position+length]; position+=length
            if code==0x000A: break
            if code==0x0200 and length>=12:
                _,last_row,_,last_col=struct.unpack_from('<IIHH',record)
                if last_row>1000 or last_col>32: raise ValueError('BIFF dimensions budget')
            if code in (0x0203,0x0204,0x00FD,0x027E,0x0205,0x0006,0x0201,0x00BD,0x00BE) and length>=4:
                row,col=struct.unpack_from('<HH',record)
                if row>=1000 or col>=32 or code in (0x00BD,0x00BE) and struct.unpack_from('<H',record,len(record)-2)[0]>=32: raise ValueError('BIFF sparse cell budget')
        sheet=book.sheet_by_index(0)
        if sheet.nrows>1000 or sheet.ncols>32: raise ValueError('spreadsheet cell budget')
        rows=[]
        for i in range(sheet.nrows):
            rows.append([xlrd.xldate_as_datetime(cell.value,book.datemode).date().isoformat() if cell.ctype==xlrd.XL_CELL_DATE else cell.value for cell in sheet.row(i)])
        return rows
    finally: book.release_resources()


def valuation(raw, cutoff):
    table=excel_rows(raw)
    headers=['日期Date','指数代码IndexCode','市盈率1（总股本）P/E1','市盈率2（计算用股本）P/E2','股息率1（总股本）D/P1','股息率2（计算用股本）D/P2']
    if not table or not all(h in table[0] for h in headers): raise ValueError('valuation columns drift')
    positions=[list(table[0]).index(h) for h in headers]; output=[]
    for cells in table[1:]:
        if not any(v is not None for v in cells): continue
        d,code,*values=[cells[i] for i in positions]
        if isinstance(d,datetime): d=d.date().isoformat()
        if type(code) in (int,float) and code==int(code): code=str(int(code)).zfill(6)
        if code!='000300': raise ValueError('valuation issuer mismatch')
        d=source_day(d)
        if d>cutoff: continue
        pe1,pe2,dy1,dy2=[num(v,lo=-1e6,hi=1e6) for v in values]
        output.append({'id':d,'date':d,'index_code':'000300','universe':'CSI300','pe_total':pe1,'pe_calculation':pe2,
                       'dividend_yield_total_pct':dy1,'dividend_yield_calculation_pct':dy2,'pe_unit':'times','dividend_unit':'percent',
                       'pe_interpretation':'nonpositive PE is loss/undefined, never cheapness rank','pb':None,'calculation_method':'provider total-share vs calculation-share bases'})
    if len({r['id'] for r in output})!=len(output): raise ValueError('conflicting valuation date')
    return sorted(output,key=lambda r:r['date'])[-60:], ['LATEST_PUBLISHED_FILE_NOT_HISTORICAL_VINTAGE']


def _quote_lines(raw):
    lines={}
    for line in raw.decode('gbk').strip().split(';'):
        if not line.strip(): continue
        match=re.fullmatch(r'\s*v_([A-Za-z0-9_.-]+)="([^"\r\n]*)"\s*',line)
        if not match or match[1] in lines: raise ValueError('quote line identity/shape')
        lines[match[1]]=match[2].split('~')
    return lines


def _local_quote_time(value, zone):
    if type(value) is not str or not re.fullmatch(r'\d{4}[-/]\d{2}[-/]\d{2} \d{2}:\d{2}:\d{2}', value): raise ValueError('quote local timestamp shape')
    normalized=value.replace('/','-')
    parsed=datetime.strptime(normalized,'%Y-%m-%d %H:%M:%S').replace(tzinfo=ZoneInfo(zone))
    if parsed.astimezone(timezone.utc).astimezone(ZoneInfo(zone)).replace(tzinfo=None)!=parsed.replace(tzinfo=None): raise ValueError('nonexistent local quote time')
    return parsed


def global_quote(raw, market, symbol, cutoff, captured):
    zone='America/New_York' if market=='US' else 'Asia/Hong_Kong' if market=='HK' else 'Asia/Shanghai'
    stock,benchmark=E.quote_keys(market,symbol)
    lines=_quote_lines(raw)
    if set(lines)!={stock,benchmark}: raise ValueError('same-response stock/benchmark identities missing')
    quotes=[]
    for key,expected in ((stock,symbol),(benchmark,'INX' if market=='US' else 'HSI' if market=='HK' else '000001')):
        row=lines[key]
        if len(row)<(73 if market=='US' else 76 if market=='HK' else 31) and key==stock or len(row)<31 or row[2]!=expected: raise ValueError('returned quote identity mismatch')
        ts_raw=row[30]
        if market=='A':
            if not re.fullmatch(r'\d{14}',ts_raw): raise ValueError('CN quote timestamp shape')
            ts_raw=datetime.strptime(ts_raw,'%Y%m%d%H%M%S').strftime('%Y-%m-%d %H:%M:%S')
        ts=_local_quote_time(ts_raw,zone)
        if ts.astimezone(timezone.utc)>C.stamp(captured): raise ValueError('future quote timestamp')
        currency=(row[35] if market=='US' else row[75] if market=='HK' else 'CNY') if key==stock else None
        expected_currency='USD' if market=='US' else 'HKD' if market=='HK' else 'CNY'
        # Benchmarks are points; vendor field must not be reused as equity price currency.
        if key==stock and currency!=expected_currency: raise ValueError('quote currency mismatch')
        price,prev,volume=num(row[3],lo=0),num(row[4],lo=0),num(row[6],lo=0)
        if price is None or price<=0 or prev is None or prev<=0: raise ValueError('missing quote price')
        quotes.append({'key':key,'symbol':expected,'name':C.text(row[1]),'price':price,'prev_close':prev,'volume':volume,
                       'source_time':ts.astimezone(timezone.utc).isoformat(timespec='microseconds').replace('+00:00','Z'),
                       'session_date':ts.date().isoformat(),'timezone':zone,'currency':expected_currency if key==stock else None,'unit':'quote currency per share' if key==stock else 'index points'})
    equity,ref=quotes; observed=C.stamp(captured).astimezone(ZoneInfo(zone)); mins=observed.hour*60+observed.minute
    phase='weekend_by_clock' if observed.weekday()>4 else 'pre_open_by_clock' if mins<570 else 'post_close_by_clock' if mins>=(960 if market=='US' else 970 if market=='HK' else 900) else 'midday_break_by_clock' if market in ('HK','A') and (720 if market=='HK' else 690)<=mins<780 else 'regular_hours_by_clock'
    stale=equity['session_date']<ref['session_date'] or (equity['volume']==0 and equity['price']==equity['prev_close'])
    return [{'id':market+':'+symbol,'market':market,'symbol':symbol,'quote':equity,'benchmark':ref,'reference_session_date':ref['session_date'],
             'clock_phase':phase,'session_calendar_confirmed':False,'source_stale':stale,
             'quote_age_seconds_at_capture':(C.stamp(captured)-C.stamp(equity['source_time'])).total_seconds(),
             'benchmark_age_seconds_at_capture':(C.stamp(captured)-C.stamp(ref['source_time'])).total_seconds(),
             'freshness_qualification':'ARCHIVED_SNAPSHOT_NOT_LIVE_FRESHNESS_CERTIFICATION','requested_cutoff':cutoff,
             'after_requested_date':equity['session_date']>cutoff or ref['session_date']>cutoff,'same_response':True}], ['CLOCK_PHASE_NOT_EXCHANGE_CALENDAR','SNAPSHOT_NOT_HISTORICAL_QUOTE']


PERIOD_CODES={'001':('annual_ytd',12),'002':('ytd',6),'003':('ytd',3),'004':('ytd',9),'006':('quarter',3),'007':('quarter',3),'008':('quarter',3)}


def financials(payload, market, symbol, cutoff, *, cn=False):
    payload=obj(payload); result=obj(payload.get('result')); rows=C.rows(result.get('data'))
    if payload.get('success') is not True: raise ValueError('datacenter response failed')
    expected_code=symbol.replace('.','_') if market=='US' else symbol
    output=[]; seen={}; counts={'fetched_rows':len(rows),'dropped_after_as_of':0,'dropped_no_notice':0,'unclassified':0,'duplicate_rows':0}
    for r in rows:
        r=obj(r); returned=r.get('SECURITY_CODE'); secucode=r.get('SECUCODE')
        expected_suffix='.HK' if market=='HK' else '.SH' if symbol.startswith(('60','68')) else '.BJ' if symbol.startswith(('43','83','87','88','92')) else '.SZ'
        if returned!=expected_code or type(secucode) is not str or (market=='US' and secucode not in [expected_code+s for s in ('.O','.N','.A')]) or (market!='US' and secucode!=expected_code+expected_suffix):
            raise ValueError('financial returned issuer mismatch')
        notice=r.get('NOTICE_DATE')
        if notice is None or notice=='': counts['dropped_no_notice']+=1; continue
        notice=source_day(notice); end=source_day(r.get('REPORT_DATE'))
        if notice>cutoff or end>cutoff: counts['dropped_after_as_of']+=1; continue
        if cn:
            if end[5:] not in ('03-31','06-30','09-30','12-31'): counts['unclassified']+=1; continue
            start=end[:4]+'-01-01'; months=int(end[5:7]); kind='annual_ytd' if months==12 else 'ytd'; code='CN_STANDARD_REPORT_END'
        else:
            try: start=source_day(r.get('START_DATE'))
            except ValueError: counts['unclassified']+=1; continue
            code=r.get('DATE_TYPE_CODE'); contract=PERIOD_CODES.get(code)
            span=(date.fromisoformat(end)-date.fromisoformat(start)).days+1
            if not contract or span<=0 or round(span/30.44)!=contract[1]: counts['unclassified']+=1; continue
            kind,months=contract
        currency='CNY' if cn else reporting_currency(r.get('CURRENCY'))
        if cn and r.get('CURRENCY') not in (None,'人民币','CNY'): raise ValueError('CN reporting currency mismatch')
        revenue=num(r.get('TOTAL_OPERATE_INCOME') if cn else r.get('OPERATE_INCOME'))
        profit=num(r.get('PARENT_NETPROFIT') if cn else r.get('PARENT_HOLDER_NETPROFIT') if market=='US' else r.get('HOLDER_PROFIT'))
        row={'id':start+'..'+end,'market':market,'symbol':symbol,'returned_secucode':secucode,'start':start,'end':end,'notice_date':notice,
             'period_kind':kind,'period_months':months,'period_days':(date.fromisoformat(end)-date.fromisoformat(start)).days+1,'period_code':code,'currency':currency,'amount_unit':'PROVIDER_RAW_SCALE_UNVERIFIED','amount_scale_to_currency_units':None,
             'revenue':revenue,'net_profit_parent':profit,'net_profit_deducted':num(r.get('DEDUCT_PARENT_NETPROFIT')) if cn else None,
             'eps_basic':num(r.get('BASIC_EPS')),'eps_diluted':num(r.get('DILUTED_EPS')),
             'eps_unit':'reporting currency per share','consolidation_scope':r.get('REPORT_TYPE') if type(r.get('REPORT_TYPE')) is str else None,
             'revision_id':None,'vintage':'CURRENT_STORED_POSTHOC','point_in_time_claim':'NONE'}
        if row['id'] in seen:
            if seen[row['id']]!=row: raise ValueError('conflicting financial revision')
            counts['duplicate_rows']+=1; continue
        seen[row['id']]=row; output.append(row)
    for r in output:
        r['missing']=[k for k in ('currency','revenue','net_profit_parent','eps_basic','consolidation_scope') if r[k] is None]
        if cn and r['net_profit_deducted'] is None: r['missing'].append('net_profit_deducted')
    output.sort(key=lambda r:(r['end'],r['start']),reverse=True)
    reasons=['CURRENT_VINTAGE_NOTICE_FILTER_ONLY']
    total=result.get('count')
    if type(total) is not int or total!=len(rows): reasons.append('PROVIDER_PAGINATION_OR_TOTAL_UNKNOWN')
    return output, reasons, counts


def _yraw(value):
    if type(value) is dict:
        if 'raw' not in value: return None
        value=value['raw']
    return num(value)


def consensus(payload, market, symbol):
    result=obj(obj(payload).get('quoteSummary'))
    if result.get('error') is not None: raise ValueError('Yahoo failure')
    result=C.rows(result.get('result'))
    if len(result)!=1: raise ValueError('Yahoo result cardinality')
    data=obj(result[0]); price=obj(data.get('price')); expected=f'{int(symbol):04d}.HK' if market=='HK' else symbol.replace('.','-')
    if price.get('symbol')!=expected: raise ValueError('Yahoo returned identity mismatch')
    currency=obj(data.get('earnings')).get('financialCurrency')
    if currency not in ('USD','HKD','CNY','EUR','GBP','JPY','CAD','AUD','CHF','TWD','KRW'): currency=None
    output=[]
    for r in C.rows(obj(data.get('earningsTrend')).get('trend')):
        r=obj(r)
        if r.get('period') not in ('0y','+1y'): continue
        end=source_day(r.get('endDate')); eps=obj(r.get('earningsEstimate')); revenue=obj(r.get('revenueEstimate'))
        avg,low,high=(_yraw(eps.get(k)) for k in ('avg','low','high')); count=_yraw(eps.get('numberOfAnalysts'))
        if count is not None and (count<0 or count!=int(count)): raise ValueError('invalid analyst count')
        if all(v is not None for v in (avg,low,high)) and not low<=avg<=high: raise ValueError('EPS bounds inverted')
        output.append({'id':r['period'],'market':market,'symbol':symbol,'returned_symbol':expected,'horizon':r['period'],'fiscal_end':end,
                       'fiscal_label':'FY'+end[:4],'period_kind':'annual_estimate','currency':currency,'quote_currency':price.get('currency'),
                       'eps_avg':avg,'eps_low':low,'eps_high':high,'eps_unit':'reporting currency per share',
                       'revenue_estimate':_yraw(revenue.get('avg')),'revenue_unit':'reporting currency units','analyst_count':count,
                       'source_update_time':None,'revision_id':None,'rating':None,
                       'missing':[k for k,v in [('currency',currency),('eps_avg',avg),('eps_low',low),('eps_high',high),('analyst_count',count)] if v is None]})
    if len({r['id'] for r in output})!=len(output): raise ValueError('duplicate consensus horizon')
    by={r['horizon']:r for r in output}
    if len(by)==2 and by['0y']['fiscal_end']>=by['+1y']['fiscal_end']: raise ValueError('consensus fiscal horizon inversion')
    return output,['PROVIDER_UPDATE_TIME_UNKNOWN','CONSENSUS_NOT_COMPANY_FACT']+([] if len(output)==2 else ['ANNUAL_HORIZON_MISSING'])


def cn_consensus(raw, symbol, cutoff):
    parser=Tables(); parser.feed(raw.decode('gbk'))
    if not re.search(r'(?<!\d)'+re.escape(symbol)+r'(?!\d)',parser.title): raise ValueError('THS returned page issuer missing')
    required=['年度','预测机构数','最小值','均值','最大值']
    tables=[t for t,context in zip(parser.tables,parser.contexts) if t and all(h in t[0] for h in required)
            and re.search(r'(?:每股收益.{0,30}元|EPS.{0,30}元[/／]股)',context,re.I)]
    if len(tables)!=1: raise ValueError('THS consensus table missing or ambiguous')
    table=tables[0]; indexes=[table[0].index(h) for h in required]; output=[]
    for cells in C.rows(table[1:]):
        if len(cells)!=len(table[0]): raise ValueError('THS table row width')
        year,count,low,avg,high=[cells[i] for i in indexes]
        if not re.fullmatch(r'20\d{2}(?:年)?',year): raise ValueError('THS fiscal year unclear')
        count=num(count,lo=0,hi=10000); low,avg,high=[num(v) for v in (low,avg,high)]
        if count is not None and count!=int(count): raise ValueError('analyst count not integer')
        if all(v is not None for v in (low,avg,high)) and not low<=avg<=high: raise ValueError('consensus bounds')
        year=year[:4]
        output.append({'id':year,'market':'A','symbol':symbol,'returned_identity_basis':'HTML page title','fiscal_end':year+'-12-31',
                       'fiscal_end_basis':'CN statutory calendar fiscal year','fiscal_label':'FY'+year,'period_kind':'annual_estimate',
                       'currency':'CNY','eps_unit':'CNY per share','eps_avg':avg,'eps_low':low,'eps_high':high,'analyst_count':count,
                       'rating':None,'source_update_time':None,'missing':[] if count and count>=3 and all(v is not None for v in (low,avg,high)) else ['EPS_OR_ANALYST_COVERAGE_LOW']})
    if len({r['id'] for r in output})!=len(output): raise ValueError('duplicate fiscal year')
    missing=[str(y) for y in range(int(cutoff[:4]),int(cutoff[:4])+3) if str(y) not in {r['id'] for r in output}]
    return output,['LATEST_SNAPSHOT_NOT_HISTORICAL_CONSENSUS','PROVIDER_UPDATE_TIME_UNKNOWN']+(['FISCAL_HORIZON_MISSING:'+','.join(missing)] if missing else [])


def normalize_extended(raw, source, subject, captured, *, now=None):
    key,cutoff,market,symbol=E.parse_subject(source,subject); spec=E.source_spec(source,subject)
    captured=C.utc(captured); now=now or datetime.now(timezone.utc)
    if type(raw) is not bytes or not 0<len(raw)<=C.MAX_BYTES or C.stamp(captured)>now: raise ValueError('invalid extended capture')
    zone='America/New_York' if market=='US' else 'Asia/Hong_Kong' if market=='HK' else 'Asia/Shanghai'
    if cutoff>C.stamp(captured).astimezone(ZoneInfo(zone)).date().isoformat(): raise ValueError('cutoff after capture')
    counters={}
    if source=='em_market_trend': rows,reasons=index_trend(C.decode(raw),key,cutoff,captured)
    elif source=='index_valuation': rows,reasons=valuation(raw,cutoff)
    elif source=='chinabond_yield_curve': rows,reasons=treasury(raw,cutoff)
    elif source in ('global_quote_ref','cn_quote_ref'): rows,reasons=global_quote(raw,market,symbol,cutoff,captured)
    elif source in ('global_financials_pit','cn_financials_ytd'): rows,reasons,counters=financials(C.decode(raw),market,symbol,cutoff,cn=source=='cn_financials_ytd')
    elif source=='yahoo_consensus': rows,reasons=consensus(C.decode(raw),market,symbol)
    else: rows,reasons=cn_consensus(raw,symbol,cutoff)
    if len(rows)>500: raise ValueError('extended normalized row budget')
    result={'schema_version':C.VERSION,'normalizer_version':'public-research-adapter.v1','source_id':source,'subject':subject,
            'market':market,'symbol':symbol,'requested_as_of':cutoff,'label':spec['label'],'status':'partial' if rows else 'missing',
            'rows':rows,'reason_codes':reasons,'counters':counters,'guard':spec['guard'],'captured_at':captured,'source_url':E.request_spec(source,subject),
            'license_status':'PROVIDER_TERMS_NOT_VERIFIED','selection_mode':'CURRENT_ARCHIVED_POSTHOC','point_in_time_claim':'NONE',
            'source_revision':None,'source_publication_time':None,'archive_age_seconds':(now-C.stamp(captured)).total_seconds(),
            'archive_stale':(now-C.stamp(captured)).total_seconds()>300 if source in ('global_quote_ref','cn_quote_ref') else (now-C.stamp(captured)).total_seconds()>86400,'coverage':'bounded provider response; missing remains unknown'}
    C.canonical(result)
    return result
